"""Deterministic ingestion helpers shared by upload and validation endpoints."""
import csv, hashlib, hmac, io, re
from collections import Counter
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

REDACTION_VERSION = "v1"
MAX_XLSX_ROWS = 100_000
MAX_XLSX_COLUMNS = 256
# 工程计划 4.1:内容长度 1—10,000 Unicode 字符,超长行记为无效
CONTENT_MAX_CHARS = 10_000
# 工程计划 4.2:标准反馈字段;映射目标只能是它们。order_id 是输入列名,
# order_ref 是存储列名,两者都接受,存回时沿用用户映射的目标名。
STANDARD_FIELDS = ('feedback_id', 'content', 'created_at', 'channel', 'product',
                   'rating', 'order_id', 'order_ref', 'status')
# 工程计划 4.2:渠道/产品缺失时落 unknown;评分只接受 1—5
DEFAULT_CHANNEL = 'unknown'
DEFAULT_PRODUCT = 'unknown'
RATING_MIN, RATING_MAX = 1, 5
# 工程计划 4.1:CSV/TXT 编码只允许 UTF-8/UTF-8 BOM 与显式 GB18030
_ENCODINGS = {
    'utf-8': 'utf-8', 'utf8': 'utf-8',
    'utf-8-sig': 'utf-8-sig', 'utf-8-bom': 'utf-8-sig',
    'gb18030': 'gb18030', 'gb-18030': 'gb18030',
}
_PATTERNS = [("email", re.compile(r"\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b")), ("phone", re.compile(r"(?<!\d)(?:\+?86[- ]?)?1[3-9]\d{9}(?!\d)")), ("order_id", re.compile(r"\b(?:ORD|ORDER)[-_]?\d{6,}\b", re.I))]
# UTF-8 BOM(U+FEFF),用 chr 写出避免源码里出现不可见字符
_BOM = chr(0xFEFF)


class UnsupportedEncoding(ValueError):
    """用户选择的编码不在允许清单内(UTF-8 / UTF-8 BOM / GB18030)。"""


class DecodeError(ValueError):
    """按所选编码无法解码上传字节。

    工程计划 4.1 禁止用替换字符悄悄兜底:解码失败必须报错并提示重新选择编码,
    否则乱码会被当成正文一路写进脱敏预览、向量与证据引文。
    """

    def __init__(self, encoding: str, error: UnicodeDecodeError):
        self.encoding = encoding
        self.start = error.start
        super().__init__(
            f'无法用 {encoding} 解码上传文件(字节偏移 {error.start});'
            '请重新选择编码:UTF-8 / UTF-8 BOM / GB18030'
        )


class MappingError(ValueError):
    """字段映射不合法(未知标准字段、重复映射、引用了不存在的源列、缺 content)。"""


def resolve_encoding(encoding: str | None) -> str:
    """把用户选择的编码名映射为 Python codec;未选择时默认 UTF-8。

    按大小写与分隔符不敏感处理:"UTF-8 BOM" / "utf_8_bom" / "utf-8-bom" 等价。
    """
    name = str(encoding or '').strip().lower().replace('_', '-').replace(' ', '-')
    if not name:
        return 'utf-8'
    if name not in _ENCODINGS:
        raise UnsupportedEncoding(
            f'不支持的编码: {name};可选: UTF-8、UTF-8 BOM、GB18030'
        )
    return _ENCODINGS[name]


def decode_text(data: bytes, encoding: str | None = None) -> str:
    """按显式编码解码上传文本,解码失败抛 DecodeError。

    解码后统一剥掉行首 BOM:无论用户选 UTF-8 还是 UTF-8 BOM,都不该让
    ``\\ufeff`` 混进第一列表头,否则列名对不上标准字段,映射整批落空。
    """
    codec = resolve_encoding(encoding)
    try:
        text = data.decode(codec)
    except UnicodeDecodeError as exc:
        raise DecodeError(codec, exc) from exc
    return text.lstrip(_BOM)

def redact_text(text: str) -> dict:
    hits = Counter(); out = text
    for kind, pattern in _PATTERNS:
        def sub(m): hits[kind] += 1; return f"<{kind.upper()}_REDACTED>"
        out = pattern.sub(sub, out)
    return {"text": out, "hits": dict(hits), "version": REDACTION_VERSION}

# 标识/定位这一行、不属于反馈正文的字段名(与 feedback.py 的取值键一致)
IDENTITY_KEYS = ('feedback_id', 'id', 'source_row')


def row_text(row) -> str:
    """把治理后的一行拼成反馈正文——逐值脱敏,并排除标识字段。

    映射后的标准行直接取 `content`(工程计划 4.2 的必填正文):created_at、
    channel 这类是元数据,拼进正文会污染向量、主题引语与规则扫描片段。
    非标准行(历史批次)保持既有口径:逐值拼接——排除 feedback_id / id /
    source_row 这些标识定位字段,它们不是用户说的话。若来源表真有一列名为
    `id` 且承载业务内容,那段内容不会被分析;这与取值逻辑一致。
    """
    content = row.get('content')
    if content is not None and str(content).strip():
        return redact_text(str(content))['text']
    return ' '.join(
        redact_text(str(value))['text']
        for key, value in row.items()
        if str(key) not in IDENTITY_KEYS and value is not None
    )


def row_source_row(row, fallback: int) -> int:
    """取行的源行号;规范化行显式带 source_row,历史行回退到枚举位置。"""
    try:
        return int(row.get('source_row'))
    except (AttributeError, TypeError, ValueError):
        return fallback


def feedback_identity(row) -> tuple[str | None, str]:
    """返回 (external_id, identity_quality)。

    工程计划 4.2:`feedback_id` 是来源系统的原始编号,存 `external_id`;平台自己的
    `id` 另生成。没有来源编号时按 `source_row` 认身份(4.4)。
    """
    external = row.get('feedback_id') if isinstance(row, dict) else None
    text = str(external).strip() if external is not None else ''
    return (text, 'source_id') if text else (None, 'source_row')


def feedback_event_key(secret: str, dataset, *, external_id: str | None, source_row: int) -> str:
    """工程计划 4.4 的事件键。

    有来源编号:`HMAC(project_dedupe_key, source_namespace + "\\0" + external_id)`
    无来源编号:`HMAC(project_dedupe_key, file_sha256 + sheet_name + source_row)`

    与 web 层的数据集级去重键**不是一回事**:那个认的是「同一个来源被重传」,
    这个认的是「同一条反馈事件」。同文本不同来源编号是两条不同事件,都要保留,
    所以消息里必须带来源身份而不是正文。

    已知缺口:XLSX 工作表选择尚未实现,`sheet_name` 目前恒为空,因此同一工作簿的
    不同工作表按行号去重会互相碰撞。补工作表选择时要连同这里一起改。
    """
    if external_id is not None:
        message = f"{dataset.get('source_namespace') or ''}\0{external_id}"
    else:
        message = f"{dataset.get('content_hash') or ''}{dataset.get('sheet_name') or ''}{source_row}"
    return hmac.new(secret.encode('utf-8'), message.encode('utf-8'), hashlib.sha256).hexdigest()


def content_fingerprint(text: str) -> str:
    """脱敏正文的 SHA-256。工程计划 4.2:只用于相似文本候选,不用于幂等。"""
    return hashlib.sha256((text or '').encode('utf-8')).hexdigest()


def build_feedback_rows(dataset, *, secret: str, redaction_version: str = REDACTION_VERSION) -> list[dict]:
    """把一个数据集的治理后行展开为待落库的 feedback 行。

    迁移回填与导入落库**共用本函数**。各写一份的话,存量行与新行的 event_key
    口径会悄悄分叉——而幂等、按行删除、跨项目约束全都建立在它之上,分叉了
    也不会有任何报错。

    规范化批次(经 `apply_mapping`)带标准字段,直接取用;历史非标准批次退回
    逐值拼接正文(`row_text`),字段只能给默认值。两种情况下正文都已经是脱敏结果。
    """
    rows = (dataset.get('preview') or {}).get('rows') or []
    source_kind = dataset.get('source_kind') or dataset.get('file_ext')
    built: list[dict] = []
    for index, row in enumerate(rows):
        if not isinstance(row, dict):
            continue
        # 与流水线、证据源查询共用同一份正文口径,否则引文 offset 在三处对不上
        text = row_text(row)
        source_row = row_source_row(row, index)
        external_id, identity_quality = feedback_identity(row)
        occurred_at = _coerce_timestamp(row.get('created_at') or row.get('occurred_at'))
        rating = _coerce_rating(row.get('rating'))
        built.append({
            'external_id': external_id,
            'event_key': feedback_event_key(secret, dataset, external_id=external_id, source_row=source_row),
            'content_redacted': text,
            'content_hash': content_fingerprint(text),
            'occurred_at': occurred_at,
            'time_quality': str(row.get('time_quality') or ('exact' if occurred_at else 'missing')),
            'channel': str(row.get('channel') or DEFAULT_CHANNEL),
            'product': str(row.get('product') or DEFAULT_PRODUCT),
            'rating': rating,
            'source_status': _optional_text(row.get('status')),
            'order_ref_redacted': _optional_text(row.get('order_ref') or row.get('order_id')),
            'source_row': source_row,
            'source_kind': source_kind,
            'identity_quality': identity_quality,
            'redaction_version': redaction_version,
            # 序号按 source_row 升序枚举,与既有 `fb_{dataset}_{序号}` 派生规则一致
            'id': f"fb_{dataset.get('id', 'ds')}_{index}",
        })
    return built


def _optional_text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return redact_text(text)['text'] if text else None


def _coerce_timestamp(value) -> datetime | None:
    """把库里存的 ISO 字符串还原成 timezone-aware datetime。

    `feedback.occurred_at` 是 timestamptz 列而不是 JSON 字符串——计划 5.1 要求
    时间筛选走数据库,`[start, end)` 半开区间才可能下推到索引。
    """
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    text = str(value or '').strip()
    if not text:
        return None
    try:
        parsed = datetime.fromisoformat(text.replace('Z', '+00:00'))
    except ValueError:
        return None
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)


def _coerce_rating(value) -> float | None:
    try:
        number = float(str(value).strip())
    except (TypeError, ValueError):
        return None
    return number if RATING_MIN <= number <= RATING_MAX else None


def run_feedback(run, repository) -> list[dict]:
    """取一个 run 输入集合内的反馈行,取自 `feedback` 实体表(工程计划 5.2)。

    「这个 run 的输入是哪几条反馈」只在这里判定:流水线、复盘、看板、导出与
    证据源查询都经过它。各写一套的话,复盘分子与主题证据会指向不同的反馈集合,
    而且不会有任何报错。

    输入取自 `run_feedbacks` 表(5.2/5.3「输入固定」),之后再导入的反馈不会隐式
    扩大这个 run 的输入。早于 0015 的 run 没有表行,回退到 run JSON 里的
    `run_feedback_ids`(0014 的过渡形态);再早于 0014 的 run 连那个字段也没有,
    只能按 dataset_ids 取——那批 run 的输入本来就没有冻结过,迁移无法凭空补出
    当时正确的集合。两级回退都是为存量数据准备的,新 run 一律走表。

    返回的行字段名与治理后行对齐(channel/product/occurred_at),所以按字段名
    筛选的调用方不需要改动。
    """
    project_id = (run or {}).get('project_id')
    run_id = (run or {}).get('id')
    frozen = repository.list_run_feedback_ids(project_id, run_id) if run_id else []
    if not frozen:
        frozen = (run or {}).get('run_feedback_ids')
    if not frozen:
        return repository.list_feedback(project_id, list((run or {}).get('dataset_ids') or []))
    # 按冻结的 id 直接取,不要拉全项目再过滤:流水线每跑一次要对同一个 run
    # 取三轮(分块、主题、风险扫描),全项目扫描会把成本乘在项目总量上
    return repository.list_feedback(project_id, feedback_ids=[str(item) for item in frozen])


def iter_run_feedback(run, repository) -> list[tuple[str, dict]]:
    """遍历一个 run 输入集合里的反馈,返回 (feedback_id, 反馈行)。"""
    return [(str(row['id']), row) for row in run_feedback(run, repository)]


def redact_row(row: dict) -> dict:
    """逐值脱敏一行。

    逐值而不是先拼成整串,避免相邻字段接出本不存在的模式(例如一列以数字结尾、
    下一列以数字开头,拼起来会被电话号码规则命中)。
    """
    return {str(key): redact_text(str(value))['text'] for key, value in row.items()}


def classify_rows(rows: list[dict], time_field: str | None = None) -> dict:
    seen = set(); valid = invalid = duplicate = redacted = missing_time = 0; preview = []
    for row in rows:
        key = tuple(sorted((str(k), str(v)) for k,v in row.items()))
        if key in seen: duplicate += 1
        seen.add(key); clean = {}; row_redacted = False
        for k,v in row.items():
            r = redact_text(str(v)); clean[k] = r['text']; row_redacted |= bool(r['hits'])
        redacted += row_redacted
        if time_field and not str(row.get(time_field, '')).strip(): missing_time += 1
        if any(str(v).strip() for v in row.values()): valid += 1
        else: invalid += 1
        if len(preview) < 5: preview.append(clean)
    return {"total": len(rows), "valid": valid, "invalid": invalid, "duplicate": duplicate, "redacted": redacted, "missing_time": missing_time, "preview": preview}

def parse_csv_text(text: str) -> dict:
    reader = csv.DictReader(io.StringIO(text, newline=''))
    headers = reader.fieldnames or []
    if not headers or any(not str(h).strip() for h in headers):
        raise ValueError('CSV header is missing or contains an empty field')
    rows = []
    for row in reader:
        line_no = reader.line_num
        if None in row:
            raise ValueError(f'CSV row {line_no} has more fields than the header')
        missing = [header for header, value in row.items() if value is None]
        if missing:
            raise ValueError(f'CSV row {line_no} has fewer fields than the header: {", ".join(missing)}')
        rows.append(row)
    # 持久业务正文只保留脱敏结果(工程计划 7.2/6.2):rows 是下游(流水线、看板、
    # 导出、证据源)唯一的取数来源,原始行不出这个函数。stats 必须基于**原始行**
    # 统计,否则「脱敏命中数」会恒为 0。
    stats = classify_rows(rows)
    return {"headers": headers, "rows": [redact_row(row) for row in rows], "stats": stats}

def parse_txt_text(text: str) -> dict:
    """解析 TXT:一行一条反馈(工程计划 4.5),空行跳过。

    返回与 CSV 相同的 {headers, rows, stats} 形状,headers 固定为 ['content'],
    这样上传预览、校验、流水线与证据源都不需要区分来源类型。与 CSV 一样,
    落库行先脱敏,stats 按**原始行**统计(脱敏命中数才不会恒为 0)。
    """
    rows = [{'content': line} for line in text.splitlines() if line.strip()]
    stats = classify_rows(rows)
    return {"headers": ['content'], "rows": [redact_row(row) for row in rows], "stats": stats}

def parse_xlsx_bytes(data: bytes) -> dict:
    """Parse the first worksheet of an XLSX workbook with bounded dimensions."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0] if wb.worksheets else None
        if ws is None: raise ValueError("workbook has no worksheets")
        values = ws.iter_rows(values_only=True)
        try: raw_headers = next(values)
        except StopIteration: return {"headers": [], "rows": [], "stats": classify_rows([])}
        headers = [str(v).strip() if v is not None else "" for v in raw_headers]
        if len(headers) > MAX_XLSX_COLUMNS: raise ValueError(f"column limit exceeded ({MAX_XLSX_COLUMNS})")
        rows = []
        for idx, vals in enumerate(values, 1):
            if idx > MAX_XLSX_ROWS: raise ValueError(f"row limit exceeded ({MAX_XLSX_ROWS})")
            row = {h: ("" if v is None else str(v)) for h, v in zip(headers, vals)}
            rows.append(row)
        # 同 CSV:落库即脱敏,stats 仍按原始行统计
        stats = classify_rows(rows)
        return {"headers": headers, "rows": [redact_row(row) for row in rows], "stats": stats}
    except ValueError: raise
    except Exception as exc:
        raise ValueError(f"invalid or corrupted xlsx file: {exc}") from exc


def validate_mapping(mapping: dict | None, headers: list[str]) -> dict[str, str]:
    """校验 源列→标准字段 映射,返回清洗后的映射。

    显式映射必须落在工程计划 4.2 的标准字段上、不能重复映射同一个目标、不能引用
    不存在的源列,且必须包含必填的 content;违反任一条抛 MappingError(端点转 422,
    不静默忽略)。未提供映射时只做同名识别(TXT 的 content、恰好同名的列),且只有
    识别出 content 才启用规范化——否则历史非标准批次会被整批判无效,那是破坏性的。
    """
    headers = [str(header) for header in headers or []]
    if not mapping:
        identity = {header: header for header in headers if header in STANDARD_FIELDS}
        return identity if 'content' in identity.values() else {}
    cleaned: dict[str, str] = {}
    targets: set[str] = set()
    for source, target in mapping.items():
        source, target = str(source), str(target or '').strip()
        if not target:
            continue
        if target not in STANDARD_FIELDS:
            raise MappingError(f'未知标准字段: {target};可选: {", ".join(STANDARD_FIELDS)}')
        if target in targets:
            raise MappingError(f'标准字段被重复映射: {target}')
        if source not in headers:
            raise MappingError(f'源列不存在于本批次: {source}')
        cleaned[source] = target
        targets.add(target)
    if not cleaned:
        identity = {header: header for header in headers if header in STANDARD_FIELDS}
        return identity if 'content' in identity.values() else {}
    if 'content' not in targets:
        raise MappingError('缺少必填字段映射: content')
    return cleaned


def _mapped_text(value, default: str) -> str:
    """取映射值并脱敏;缺失/空白落默认值(工程计划 4.2:unknown)。"""
    text = str(value).strip() if value is not None else ''
    return redact_text(text)['text'] if text else default


def _parse_timestamp(value, tzinfo) -> tuple[str | None, str | None]:
    """解析时间戳,返回 (ISO 字符串, 错误码);错误码 ∈ invalid/nonexistent/ambiguous。

    工程计划 4.5:带偏移时间直接转 UTC;无偏移时间按用户选定时区解释;夏令时的
    歧义(fall back)与不存在(spring forward)必须提示,不能由库静默选一个结果。
    """
    text = str(value or '').strip().replace('/', '-')
    if not text:
        return None, 'invalid'
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        return None, 'invalid'
    if parsed.tzinfo is not None:
        return parsed.astimezone(timezone.utc).isoformat(), None
    fold0 = parsed.replace(tzinfo=tzinfo, fold=0)
    fold1 = parsed.replace(tzinfo=tzinfo, fold=1)
    if fold0.utcoffset() != fold1.utcoffset():
        back = fold0.astimezone(timezone.utc).astimezone(tzinfo).replace(tzinfo=None)
        if back != parsed:
            return None, 'nonexistent'
        return None, 'ambiguous'
    return fold0.isoformat(), None


def apply_mapping(rows: list[dict], mapping: dict[str, str], *,
                  time_policy: str = 'static', timezone_name: str = 'Asia/Shanghai') -> dict:
    """按映射把源行规范化为工程计划 4.2 的标准字段行。

    约束:
    - content 必填,缺失/纯空白/超长(>10,000 Unicode 字符)或公式字符串都记为无效;
      公式字符串不能被当成客诉正文(工程计划 4.3),应提示用户导出为值。
    - rating 只接受配置的 1—5 数值,越界或非数值无效。
    - created_at:严格模式必填且必须可解析;静态模式缺失时行仍有效并标
      time_quality=missing(工程计划 4.2);无偏移时间按 timezone_name 解释。
    - channel/product 缺失落 unknown。

    无效行不落库,但 health.errors 保留 source_row 与错误码,用户可按行号修正;
    有效行带 source_row 与 time_quality。返回 {"rows","stats","errors"},stats 满足
    input = valid + invalid + duplicate(工程计划 4.5)。
    """
    tzinfo = ZoneInfo(timezone_name)
    by_target = {target: source for source, target in mapping.items()}
    stored: list[dict] = []
    errors: list[dict] = []
    seen: set[tuple] = set()
    invalid = duplicate = missing_time = 0

    def value(row: dict, field: str):
        source = by_target.get(field)
        return row.get(source) if source is not None else None

    for index, row in enumerate(rows):
        if not isinstance(row, dict):
            invalid += 1
            errors.append({'source_row': index, 'code': 'INVALID_ROW', 'field': None})
            continue
        row_errors: list[tuple[str | None, str]] = []
        raw_content = value(row, 'content')
        content = str(raw_content).strip() if raw_content is not None else ''
        if content.startswith('='):
            row_errors.append(('content', 'FORMULA_NOT_ALLOWED'))
        elif not content:
            row_errors.append(('content', 'CONTENT_REQUIRED'))
        elif len(content) > CONTENT_MAX_CHARS:
            row_errors.append(('content', 'CONTENT_TOO_LONG'))

        occurred_at = None
        time_quality = 'missing'
        raw_time = value(row, 'created_at')
        if raw_time is not None and str(raw_time).strip():
            occurred_at, parse_error = _parse_timestamp(raw_time, tzinfo)
            if parse_error is None:
                time_quality = 'exact'
            elif time_policy == 'strict':
                row_errors.append(('created_at', {
                    'invalid': 'INVALID_CREATED_AT',
                    'nonexistent': 'TIMEZONE_NONEXISTENT',
                    'ambiguous': 'TIMEZONE_AMBIGUOUS',
                }[parse_error]))
        elif time_policy == 'strict':
            row_errors.append(('created_at', 'CREATED_AT_REQUIRED'))

        rating = None
        raw_rating = value(row, 'rating')
        if raw_rating is not None and str(raw_rating).strip():
            try:
                number = float(str(raw_rating).strip())
            except ValueError:
                row_errors.append(('rating', 'RATING_OUT_OF_RANGE'))
            else:
                if not RATING_MIN <= number <= RATING_MAX:
                    row_errors.append(('rating', 'RATING_OUT_OF_RANGE'))
                else:
                    rating = int(number) if number.is_integer() else number

        if row_errors:
            invalid += 1
            errors.extend({'source_row': index, 'code': code, 'field': field}
                          for field, code in row_errors)
            continue

        normalised = {
            'source_row': index,
            'content': redact_text(content)['text'],
            'channel': _mapped_text(value(row, 'channel'), DEFAULT_CHANNEL),
            'product': _mapped_text(value(row, 'product'), DEFAULT_PRODUCT),
            'time_quality': time_quality,
        }
        if occurred_at is not None:
            normalised['created_at'] = occurred_at
        if rating is not None:
            normalised['rating'] = rating
        for field in ('feedback_id', 'order_id', 'order_ref', 'status'):
            kept = value(row, field)
            if kept is not None and str(kept).strip():
                normalised[field] = redact_text(str(kept).strip())['text']
        key = tuple(sorted((k, str(v)) for k, v in normalised.items() if k != 'source_row'))
        if key in seen:
            duplicate += 1
            continue
        seen.add(key)
        if time_quality == 'missing':
            missing_time += 1
        stored.append(normalised)

    return {
        'rows': stored,
        'stats': {
            'total': len(rows),
            'valid': len(stored),
            'invalid': invalid,
            'duplicate': duplicate,
            # 命中数由解析函数按原始行统计;调用方在 validate 端点合并
            'redacted': 0,
            'missing_time': missing_time,
        },
        'errors': errors,
    }
